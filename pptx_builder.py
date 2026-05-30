"""Prezentatsiya yig'uvchi — python-pptx orqali chiroyli slaydlar yasaydi."""
import io
import os
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

import ai_images

# Rang sxemasi
BG = RGBColor(0x0F, 0x17, 0x2A)        # quyuq ko'k
ACCENT = RGBColor(0x38, 0xBD, 0xF8)    # och ko'k
TEXT = RGBColor(0xE6, 0xED, 0xF3)      # oqimtir
SUBTLE = RGBColor(0x94, 0xA3, 0xB8)    # kulrang

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _solid_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return tf


def build_presentation(data: dict, image_source: str, topic: str) -> str:
    """image_source: 'ai' yoki 'web'. Tayyor .pptx fayl yo'lini qaytaradi."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ---- Title slide ----
    s = prs.slides.add_slide(blank)
    _solid_bg(s, BG)
    # accent chiziq
    bar = s.shapes.add_shape(1, Inches(0.8), Inches(3.2), Inches(2.2), Inches(0.12))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

    tf = _add_textbox(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.5))
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = data.get("title", topic)
    r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = TEXT
    if data.get("subtitle"):
        p2 = tf.add_paragraph()
        r2 = p2.add_run(); r2.text = data["subtitle"]
        r2.font.size = Pt(20); r2.font.color.rgb = SUBTLE

    # ---- Content slides ----
    slides = data.get("slides", [])
    for s_data in slides:
        _add_content_slide(prs, blank, s_data, image_source)

    out_dir = tempfile.mkdtemp(prefix="slidebot_")
    safe = "".join(c for c in topic if c.isalnum() or c in " _-").strip()[:40] or "presentation"
    path = os.path.join(out_dir, f"{safe}.pptx")
    prs.save(path)
    return path


def _add_content_slide(prs, layout, s_data, image_source):
    slide = prs.slides.add_slide(layout)
    _solid_bg(slide, BG)

    img_bytes = _get_image(s_data, image_source)
    has_img = img_bytes is not None

    text_width = Inches(7.2) if has_img else Inches(11.7)

    # Sarlavha
    title_tf = _add_textbox(slide, Inches(0.8), Inches(0.6), text_width, Inches(1.2))
    tp = title_tf.paragraphs[0]
    tr = tp.add_run(); tr.text = s_data["title"]
    tr.font.size = Pt(30); tr.font.bold = True; tr.font.color.rgb = ACCENT

    # Bulletlar
    body_tf = _add_textbox(slide, Inches(0.8), Inches(1.9), text_width, Inches(4.8))
    first = True
    for b in s_data["bullets"]:
        p = body_tf.paragraphs[0] if first else body_tf.add_paragraph()
        first = False
        run = p.add_run(); run.text = f"•  {b}"
        run.font.size = Pt(18); run.font.color.rgb = TEXT
        p.space_after = Pt(10)
        p.alignment = PP_ALIGN.LEFT

    # Rasm
    if has_img:
        png = ai_images.normalize_to_png(img_bytes)
        if png:
            try:
                slide.shapes.add_picture(
                    io.BytesIO(png), Inches(8.4), Inches(1.6),
                    width=Inches(4.2), height=Inches(4.2)
                )
            except Exception as e:
                print(f"[pptx] rasm qo'shilmadi: {e}")


def _get_image(s_data, image_source) -> bytes | None:
    if image_source == "ai":
        return ai_images.generate_ai_image(s_data.get("image_prompt", s_data["title"]))
    elif image_source == "web":
        return ai_images.fetch_web_image(s_data.get("image_query", s_data["title"]))
    return None
